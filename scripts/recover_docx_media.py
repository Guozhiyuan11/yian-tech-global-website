from __future__ import annotations

import argparse
import io
import json
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps


REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DRAW_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def read_zip_member(archive: zipfile.ZipFile, name: str) -> tuple[bytes, bool]:
    """Read a ZIP member and recover valid payload bytes when only the stored CRC is wrong."""
    try:
        return archive.read(name), False
    except zipfile.BadZipFile:
        stream = archive.open(name)
        if hasattr(stream, "_expected_crc"):
            stream._expected_crc = None  # type: ignore[attr-defined]
        return stream.read(), True


def decode_image(data: bytes) -> Image.Image:
    image = Image.open(io.BytesIO(data))
    image.load()
    return image


def normalised_pixels(data: bytes, size: int = 24) -> np.ndarray:
    image = decode_image(data).convert("RGB")
    contained = ImageOps.contain(image, (size, size), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (size, size), tuple(ImageStat.mean(image)))
    x = (size - contained.width) // 2
    y = (size - contained.height) // 2
    canvas.paste(contained, (x, y))
    return np.asarray(canvas, dtype=np.float32) / 255.0


class ImageStat:
    @staticmethod
    def mean(image: Image.Image) -> tuple[int, int, int]:
        sample = image.convert("RGB").resize((1, 1), Image.Resampling.BOX)
        return tuple(int(value) for value in sample.getpixel((0, 0)))


def ordered_docx_media(archive: zipfile.ZipFile) -> list[str]:
    document_xml, _ = read_zip_member(archive, "word/document.xml")
    rels_xml, _ = read_zip_member(archive, "word/_rels/document.xml.rels")
    root = ElementTree.fromstring(document_xml)
    rel_root = ElementTree.fromstring(rels_xml)
    rels = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rel_root.findall(f"{{{REL_NS}}}Relationship")
        if rel.attrib.get("Type", "").endswith("/image")
    }

    ordered: list[str] = []
    for blip in root.iter(f"{{{DRAW_NS}}}blip"):
        rel_id = blip.attrib.get(f"{{{OFFICE_REL_NS}}}embed")
        target = rels.get(rel_id or "")
        if not target:
            continue
        member = str(PurePosixPath("word") / PurePosixPath(target))
        if member not in ordered:
            ordered.append(member)
    return ordered


def workbook_thumbnails(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    image_sheet = next(sheet for sheet in workbook.worksheets if len(getattr(sheet, "_images", [])) == 40)
    items = []
    for image in sorted(image_sheet._images, key=lambda item: item.anchor._from.row):
        row = image.anchor._from.row + 1
        source_id = int(image_sheet.cell(row=row, column=2).value)
        data = image._data()
        items.append({"source_id": source_id, "data": data, "pixels": normalised_pixels(data)})
    return items


def create_contact_sheets(records: list[dict], output_dir: Path) -> list[str]:
    contact_dir = output_dir / "contact-sheets"
    contact_dir.mkdir(parents=True, exist_ok=True)
    font = ImageFont.load_default(size=18)
    paths = []
    columns = 5
    rows = 2
    tile_width = 320
    tile_height = 250
    image_height = 210

    for page, start in enumerate(range(0, len(records), columns * rows), start=1):
        page_records = records[start : start + columns * rows]
        sheet = Image.new("RGB", (columns * tile_width, rows * tile_height), "#111311")
        draw = ImageDraw.Draw(sheet)
        for index, record in enumerate(page_records):
            row, column = divmod(index, columns)
            left = column * tile_width
            top = row * tile_height
            image = Image.open(record["saved_path"]).convert("RGB")
            preview = ImageOps.contain(image, (tile_width - 12, image_height - 12), Image.Resampling.LANCZOS)
            image_box = Image.new("RGB", (tile_width - 12, image_height - 12), "#d8d7d0")
            px = (image_box.width - preview.width) // 2
            py = (image_box.height - preview.height) // 2
            image_box.paste(preview, (px, py))
            sheet.paste(image_box, (left + 6, top + 6))
            label = f"{record['source_id']:02d}  {record['width']}x{record['height']}  {record['format']}"
            draw.text((left + 8, top + image_height + 8), label, font=font, fill="#d8ff45")
        contact_path = contact_dir / f"contact-{page:02d}.jpg"
        sheet.save(contact_path, quality=92)
        paths.append(str(contact_path))
    return paths


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    thumbnails = workbook_thumbnails(args.xlsx)
    thumbnail_matrix = np.stack([item["pixels"] for item in thumbnails])

    records = []
    with zipfile.ZipFile(args.docx) as archive:
        media_names = ordered_docx_media(archive)
        for doc_order, member in enumerate(media_names, start=1):
            data, crc_ignored = read_zip_member(archive, member)
            image = decode_image(data)
            source_format = (image.format or Path(member).suffix.lstrip(".") or "JPEG").upper()
            extension = ".png" if source_format == "PNG" else ".jpg"

            pixels = normalised_pixels(data)
            distances = np.mean((thumbnail_matrix - pixels) ** 2, axis=(1, 2, 3))
            nearest_index = int(np.argmin(distances))
            mapped_source_id = thumbnails[nearest_index]["source_id"]
            second_distance = float(np.partition(distances, 1)[1])

            output_path = args.output_dir / f"source-{mapped_source_id:02d}{extension}"
            if extension == ".png":
                image.save(output_path, format="PNG", optimize=True)
            else:
                image.convert("RGB").save(output_path, format="JPEG", quality=95, optimize=True)

            records.append(
                {
                    "source_id": mapped_source_id,
                    "doc_order": doc_order,
                    "docx_member": member,
                    "crc_ignored": crc_ignored,
                    "format": source_format,
                    "width": image.width,
                    "height": image.height,
                    "match_distance": float(distances[nearest_index]),
                    "second_distance": second_distance,
                    "saved_path": str(output_path),
                }
            )

    mapped_ids = [record["source_id"] for record in records]
    if sorted(mapped_ids) != list(range(1, 41)):
        raise RuntimeError(f"Image-to-row matching was not one-to-one: {sorted(mapped_ids)}")

    records.sort(key=lambda record: record["source_id"])
    contact_sheets = create_contact_sheets(records, args.output_dir)
    report = {
        "source_docx": str(args.docx),
        "source_xlsx": str(args.xlsx),
        "record_count": len(records),
        "mapped_ids": mapped_ids,
        "records": records,
        "contact_sheets": contact_sheets,
    }
    report_path = args.output_dir / "recovery-report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(
        json.dumps(
            {
                "record_count": report["record_count"],
                "unique_mapped_ids": len(set(mapped_ids)),
                "crc_ignored_count": sum(record["crc_ignored"] for record in records),
                "min_width": min(record["width"] for record in records),
                "max_width": max(record["width"] for record in records),
                "min_height": min(record["height"] for record in records),
                "max_height": max(record["height"] for record in records),
                "contact_sheets": contact_sheets,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
