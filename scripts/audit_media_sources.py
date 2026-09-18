from __future__ import annotations

import argparse
import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook
from PIL import Image


def image_size(data: bytes) -> dict[str, Any]:
    try:
        with Image.open(io.BytesIO(data)) as image:
            return {
                "format": image.format,
                "width": image.width,
                "height": image.height,
                "mode": image.mode,
            }
    except Exception as exc:  # pragma: no cover - source diagnostics
        return {"error": str(exc)}


def decode_csv(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", raw, 0, 1, "Unsupported encoding")


def audit_csv(path: Path) -> dict[str, Any]:
    text, encoding = decode_csv(path)
    rows = list(csv.DictReader(io.StringIO(text)))
    return {
        "path": str(path),
        "encoding": encoding,
        "headers": list(rows[0].keys()) if rows else [],
        "row_count": len(rows),
        "rows": rows,
    }


def audit_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = [
            [cell.value for cell in row]
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column)
        ]
        images = []
        for order, item in enumerate(getattr(sheet, "_images", []), start=1):
            anchor = getattr(item, "anchor", None)
            marker = getattr(anchor, "_from", None)
            data = item._data()
            images.append(
                {
                    "order": order,
                    "anchor_row_1based": marker.row + 1 if marker is not None else None,
                    "anchor_col_1based": marker.col + 1 if marker is not None else None,
                    "declared_width": item.width,
                    "declared_height": item.height,
                    "filename": getattr(item, "path", None),
                    **image_size(data),
                }
            )
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
                "images": images,
            }
        )
    return {"path": str(path), "sheet_count": len(sheets), "sheets": sheets}


def audit_docx(path: Path) -> dict[str, Any]:
    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    inline_shape_count: int | None = None
    parser_error: str | None = None

    try:
        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
        inline_shape_count = len(document.inline_shapes)
    except Exception as exc:  # continue with raw OOXML when a media CRC is bad
        parser_error = f"{type(exc).__name__}: {exc}"

        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        for paragraph in root.iter(f"{word_ns}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{word_ns}t")).strip()
            if text:
                paragraphs.append(text)
        for table in root.iter(f"{word_ns}tbl"):
            table_rows = []
            for row in table.findall(f"{word_ns}tr"):
                cells = []
                for cell in row.findall(f"{word_ns}tc"):
                    text = "".join(node.text or "" for node in cell.iter(f"{word_ns}t")).strip()
                    cells.append(text)
                table_rows.append(cells)
            tables.append(table_rows)

    media = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(item for item in archive.namelist() if item.startswith("word/media/")):
            try:
                data = archive.read(name)
                media.append({"name": name, "bytes": len(data), "readable": True, **image_size(data)})
            except Exception as exc:
                media.append({"name": name, "readable": False, "error": f"{type(exc).__name__}: {exc}"})

    return {
        "path": str(path),
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "table_count": len(tables),
        "tables": tables,
        "inline_shape_count": inline_shape_count,
        "parser_error": parser_error,
        "media": media,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--csv", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    result = {
        "docx": audit_docx(args.docx),
        "xlsx": audit_xlsx(args.xlsx),
        "csv": audit_csv(args.csv),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(
        json.dumps(
            {
                "docx": {
                    "paragraphs": result["docx"]["paragraph_count"],
                    "tables": result["docx"]["table_count"],
                    "inline_shapes": result["docx"]["inline_shape_count"],
                    "media": len(result["docx"]["media"]),
                },
                "xlsx": [
                    {
                        "sheet": sheet["title"],
                        "rows": sheet["max_row"],
                        "columns": sheet["max_column"],
                        "images": len(sheet["images"]),
                    }
                    for sheet in result["xlsx"]["sheets"]
                ],
                "csv": {
                    "encoding": result["csv"]["encoding"],
                    "rows": result["csv"]["row_count"],
                    "headers": result["csv"]["headers"],
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
